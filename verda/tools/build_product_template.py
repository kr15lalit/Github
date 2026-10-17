"""
Build verda/Verda-Product-Catalogue-Template.xlsx: a short fill-in sheet for the product team.

Run: python3 verda/tools/build_product_template.py
"""
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "Verda-Product-Catalogue-Template.xlsx")

F = "Arial"
GREEN, DARK, GOLD = "1F5F23", "13261A", "9C6B12"
FILL_HEAD = PatternFill("solid", fgColor=GREEN)
FILL_INPUT = PatternFill("solid", fgColor="FFF8DC")
FILL_AUTO = PatternFill("solid", fgColor="E8EFE8")
FILL_SAMPLE = PatternFill("solid", fgColor="FFFDF2")
THIN = Side(style="thin", color="D5D9D5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CATEGORIES = ["Healthy Snacks", "Makhana", "Dry Fruits & Nuts", "Diwali Gift Hampers", "Atta, Rice & Dals", "Oils & Ghee", "Millets"]

# key, header, width, help
COLUMNS = [
    ("name", "Product Name", 30, "Same name on several rows = one product with a size/flavour choice on the website."),
    ("category", "Category", 20, "Pick from the dropdown."),
    ("flavour", "Flavour / Type\n(optional)", 18, "Leave blank if there is only one."),
    ("size", "Pack Size", 11, "As printed on the pack, e.g. 100 g, 1 kg, 500 ml, 1 L."),
    ("mrp", "MRP (₹)", 10, "Number only, no ₹ sign."),
    ("price", "Selling Price (₹)", 12, "Price on the website. Same as MRP if no discount. Cannot be more than MRP."),
    ("disc", "Discount", 10, "Calculated automatically."),
    ("stock", "Stock", 9, "Packs available now."),
    ("short", "Short Description (1–2 lines)", 44, "What makes it special, in one or two lines."),
    ("photo", "Front Photo\n(file name)", 26, "Name of the product photo file you will send."),
    ("back", "Back-of-Pack Photo\n(file name)", 30, "Photo of the back label. We take ingredients, nutrition, FSSAI no. and shelf life from it."),
    ("notes", "Notes\n(optional)", 26, "Anything else, e.g. hamper contents."),
]

SAMPLES = [
    ("Jowar Puffs", "Healthy Snacks", "Classic Salted", "100 g", 99, 89, 120, "Light, crunchy jowar puffs with a pinch of salt. Roasted, never fried.", "jowar-classic-100g.jpg", "jowar-classic-100g-back.jpg", ""),
    ("Jowar Puffs", "Healthy Snacks", "Classic Salted", "200 g", 179, 159, 80, "Family pack of our classic salted jowar puffs.", "jowar-classic-200g.jpg", "jowar-classic-200g-back.jpg", "Same name = shown as a size option"),
    ("Jowar Puffs", "Healthy Snacks", "Peri Peri", "100 g", 99, 89, 100, "Jowar puffs tossed in a fiery peri peri seasoning.", "jowar-periperi-100g.jpg", "jowar-periperi-100g-back.jpg", "Same name = shown as a flavour option"),
    ("Premium Phool Makhana", "Makhana", "", "250 g", 399, 349, 90, "Hand-picked, large-size fox nuts from Bihar.", "makhana-250g.jpg", "makhana-250g-back.jpg", ""),
    ("California Almonds", "Dry Fruits & Nuts", "", "500 g", 649, 579, 90, "Crunchy, naturally sweet California almonds.", "almonds-500g.jpg", "almonds-500g-back.jpg", ""),
    ("Diwali Delight Gift Hamper", "Diwali Gift Hampers", "", "1 box", 1499, 1299, 40, "Almonds, cashews, raisins and roasted makhana in a festive keepsake box.", "hamper-delight.jpg", "", "Contents: almonds 200 g, cashews 200 g, raisins 150 g, makhana 80 g"),
    ("Cold Pressed Mustard Oil", "Oils & Ghee", "", "1 L", 249, 249, 90, "Pure kachi ghani mustard oil, cold pressed.", "mustard-oil-1l.jpg", "mustard-oil-1l-back.jpg", "No discount: price = MRP"),
]


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.sheet_view.showGridLines = False

    # Short instructions on the same sheet.
    ws["A1"] = "Verda — Product List"
    ws["A1"].font = Font(name=F, bold=True, size=16, color=GREEN)
    notes = [
        "• One row per pack size or flavour. Rows with the same Product Name become one product with options.",
        "• Fill the yellow cells. Rows 7–13 are examples — type over them or delete them.",
        "• Send the filled sheet back with all photos (front + back of pack) in one folder or zip.",
    ]
    for i, t in enumerate(notes, start=2):
        ws.cell(i, 1, t).font = Font(name=F, size=10, color="444444")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(COLUMNS))

    head = 6
    col = {}
    for ci, (key, header, width, help_text) in enumerate(COLUMNS, start=1):
        col[key] = get_column_letter(ci)
        c = ws.cell(head, ci, header)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
        c.comment = Comment(help_text, "Verda")
        ws.column_dimensions[col[key]].width = width
    ws.row_dimensions[head].height = 34

    first, last = head + 1, head + 200
    for r in range(first, last + 1):
        sample = SAMPLES[r - first] if r - first < len(SAMPLES) else None
        values = dict(zip([k for k, *_ in COLUMNS if k != "disc"], sample)) if sample else {}
        for ci, (key, *_rest) in enumerate(COLUMNS, start=1):
            c = ws.cell(r, ci)
            c.border = BORDER
            c.font = Font(name=F, size=10, italic=bool(sample), color="333333")
            c.alignment = Alignment(vertical="top", wrap_text=key in ("short", "notes"))
            if key == "disc":
                m, p = col["mrp"] + str(r), col["price"] + str(r)
                c.value = '=IF(AND(ISNUMBER(%s),ISNUMBER(%s),%s>0),(%s-%s)/%s,"")' % (m, p, m, m, p, m)
                c.number_format = "0%;-0%;-"
                c.fill = FILL_AUTO
                continue
            c.fill = FILL_SAMPLE if sample else FILL_INPUT
            if values.get(key) not in (None, ""):
                c.value = values[key]
            if key in ("mrp", "price"):
                c.number_format = "₹#,##0"

    def dv(rule, key):
        ws.add_data_validation(rule)
        rule.add("%s%d:%s%d" % (col[key], first, col[key], last))

    dv(DataValidation(type="list", formula1="Categories!$A$2:$A$31", allow_blank=True, showErrorMessage=True, error="Pick a category, or add a new one on the Categories sheet."), "category")
    for key in ("mrp", "stock"):
        dv(DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, error="Numbers only, please."), key)
    p, m = col["price"], col["mrp"]
    dv(DataValidation(type="custom", formula1='OR(%s%d="",%s%d="",AND(ISNUMBER(%s%d),%s%d<=%s%d))' % (p, first, m, first, p, first, p, first, m, first), allow_blank=True, showErrorMessage=True, error="Selling Price must be a number and not more than MRP."), "price")

    ws.freeze_panes = ws.cell(first, 2)

    # Categories: drives the dropdown; new ones can be added at the bottom.
    wc = wb.create_sheet("Categories")
    wc.sheet_view.showGridLines = False
    wc["A1"] = "Category (add new ones below)"
    wc["A1"].font = Font(name=F, bold=True, color="FFFFFF")
    wc["A1"].fill = FILL_HEAD
    wc.column_dimensions["A"].width = 32
    for i in range(30):
        c = wc.cell(i + 2, 1, CATEGORIES[i] if i < len(CATEGORIES) else None)
        c.font = Font(name=F, size=10)
        c.fill = FILL_INPUT
        c.border = BORDER

    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
